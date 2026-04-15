# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl", "click", "pymupdf"]
# ///

from pathlib import Path
import re
import xml.etree.ElementTree as ET

import click
import openpyxl


@click.group()
def cli():
    pass


@cli.command()
@click.argument("input", type=click.Path(exists=True, path_type=Path))
@click.argument("output", type=click.Path(path_type=Path))
def split(input, output):
    """Split an Excel workbook into separate .xlsx files, one per worksheet."""
    output.mkdir(parents=True, exist_ok=True)
    source = openpyxl.load_workbook(input)

    for sheet_name in source.sheetnames:
        wb = openpyxl.load_workbook(input)
        for name in wb.sheetnames:
            if name != sheet_name:
                del wb[name]
        out_path = output / sheet_name / "format.xlsx"
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
        click.echo(out_path, err=True)

    source.close()


@cli.command()
@click.argument("input", type=click.Path(exists=True, path_type=Path))
def extract_text(input):
    """Extract raw text from a PDF and save as *.raw.txt."""
    import pymupdf

    doc = pymupdf.open(input)
    text = "\n\n".join(page.get_text() for page in doc)
    doc.close()

    out_path = input.with_suffix(".raw.txt")
    out_path.write_text(text)
    click.echo(out_path, err=True)


@cli.command()
@click.argument("input", type=click.Path(exists=True, path_type=Path))
def clean_text(input):
    """Clean up raw extracted PDF text into readable paragraphs."""
    import re

    raw = Path(input).read_text()

    # Remove soft hyphens at line breaks (join the word)
    text = re.sub(r"\xad\s*\n\s*", "", raw)

    # Remove repeated header/footer lines
    text = re.sub(
        r"^.*INSTRUCTIONS FOR FEC FORM \S+ AND RELATED SCHEDULES.*$",
        "",
        text,
        flags=re.MULTILINE,
    )
    text = re.sub(
        r"^\s*Federal Election Commission \(Revised.*$",
        "",
        text,
        flags=re.MULTILINE,
    )
    text = re.sub(r"^\s*Page \d+\s*$", "", text, flags=re.MULTILINE)

    # Reflow paragraphs: join short lines that are continuations
    lines = text.split("\n")
    reflowed = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            reflowed.append("")
            continue

        # Check if this line continues the previous paragraph:
        # previous line exists, is not empty, doesn't end with a colon,
        # and this line starts with lowercase or continues a sentence
        if (
            reflowed
            and reflowed[-1]
            and not reflowed[-1].endswith(":")
            and stripped[0].islower()
        ):
            reflowed[-1] = reflowed[-1] + " " + stripped
        else:
            reflowed.append(stripped)

    # Collapse runs of blank lines
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(reflowed)).strip()

    out_path = Path(input).with_suffix("").with_suffix(".txt")
    out_path.write_text(text + "\n")
    click.echo(out_path, err=True)


@cli.command()
@click.argument("input", type=click.Path(exists=True, path_type=Path))
@click.option("--page", "page_number", type=int, required=True, help="1-based PDF page number.")
@click.option(
    "--contains",
    help="Only include rows whose label contains this text, case-insensitive.",
)
def extract_boxes(input, page_number, contains):
    """Extract inferred input box locations from a non-fillable PDF page."""
    import json
    import re
    import subprocess
    import tempfile

    if page_number < 1:
        raise click.ClickException("--page must be >= 1")

    with tempfile.TemporaryDirectory() as tmpdir:
        bbox_path = Path(tmpdir) / "page.bbox.html"
        svg_path = Path(tmpdir) / "page.svg"

        subprocess.run(
            [
                "pdftotext",
                "-f",
                str(page_number),
                "-l",
                str(page_number),
                "-bbox-layout",
                str(input),
                str(bbox_path),
            ],
            check=True,
        )
        subprocess.run(
            [
                "pdftocairo",
                "-svg",
                "-f",
                str(page_number),
                "-l",
                str(page_number),
                str(input),
                str(svg_path),
            ],
            check=True,
        )

        labels = _extract_label_rows_from_bbox(bbox_path)
        boxes = _extract_amount_boxes_from_svg(svg_path)
        matches = _match_labels_to_boxes(labels, boxes, page_number)

    if contains:
        needle = contains.lower()
        matches = [row for row in matches if needle in row["label"].lower()]

    click.echo(json.dumps(matches, indent=2))


def _extract_label_rows_from_bbox(path: Path) -> list[dict]:
    tree = ET.parse(path)
    root = tree.getroot()
    page = root.find(".//page")
    if page is None:
        return []

    rows = []
    current = None
    continuation_prefixes = ("(ii)", "(iii)", "(iv)", "(b)", "(c)", "(d)", "(e)")

    for line in page.findall(".//line"):
        words = line.findall("word")
        if not words:
            continue

        text = " ".join((word.text or "").strip() for word in words).strip()
        if not text:
            continue

        y = float(line.attrib["yMin"])
        x = float(line.attrib["xMin"])

        if x >= 220:
            continue

        if re.match(r"^\d+\.", text) or text.startswith("(a)"):
            current = {
                "label": text,
                "label_y": y,
            }
            rows.append(current)
            continue

        if current is None:
            continue

        if x <= 90 and text.startswith(continuation_prefixes):
            current["label"] += " " + text
            current["label_y"] = min(current["label_y"], y)
            continue

        if x <= 90 and text.startswith("("):
            current = None
            continue

        if x <= 90 and current["label"].startswith("(a)"):
            current["label"] += " " + text
            current["label_y"] = min(current["label_y"], y)
            continue

    return rows


def _extract_amount_boxes_from_svg(path: Path) -> list[dict]:
    svg = path.read_text()
    pattern = re.compile(
        r'<path fill="none" stroke-width="2\.5"[^>]*'
        r'd="M ([0-9.]+) ([0-9.]+) L ([0-9.]+) [0-9.]+ L [0-9.]+ ([0-9.]+) [^"]*"'
    )

    boxes = []
    for match in pattern.finditer(svg):
        x0 = float(match.group(1))
        y_top_pdf = float(match.group(2))
        x1 = float(match.group(3))
        y_bottom_pdf = float(match.group(4))

        width = round(x1 - x0, 3)
        height = round(y_top_pdf - y_bottom_pdf, 3)

        if not (150 <= width <= 170 and 18 <= height <= 20):
            continue

        boxes.append(
            {
                "x0": round(x0, 3),
                "y0_pdf": round(y_bottom_pdf, 3),
                "x1": round(x1, 3),
                "y1_pdf": round(y_top_pdf, 3),
                "width": width,
                "height": height,
                "y_top": round(792 - y_top_pdf, 3),
            }
        )

    boxes.sort(key=lambda box: (box["y_top"], box["x0"]))
    return boxes


def _match_labels_to_boxes(
    labels: list[dict], boxes: list[dict], page_number: int
) -> list[dict]:
    rows_by_y: dict[float, list[dict]] = {}
    for box in boxes:
        rows_by_y.setdefault(box["y_top"], []).append(box)

    row_centers = sorted(rows_by_y)
    matches = []

    for label in labels:
        eligible_rows = [y for y in row_centers if y >= label["label_y"]]
        if not eligible_rows:
            continue

        row_y = eligible_rows[0]
        row_boxes = sorted(rows_by_y[row_y], key=lambda box: box["x0"])
        if len(row_boxes) < 2:
            continue

        column_a, column_b = row_boxes[0], row_boxes[1]
        matches.append(
            {
                "label": label["label"],
                "page": page_number,
                "label_y_top": round(label["label_y"], 3),
                "column_a": {
                    "top_left": {
                        "x": column_a["x0"],
                        "y": column_a["y_top"],
                    },
                    "bbox_pdf": [
                        column_a["x0"],
                        column_a["y0_pdf"],
                        column_a["x1"],
                        column_a["y1_pdf"],
                    ],
                    "width": column_a["width"],
                    "height": column_a["height"],
                },
                "column_b": {
                    "top_left": {
                        "x": column_b["x0"],
                        "y": column_b["y_top"],
                    },
                    "bbox_pdf": [
                        column_b["x0"],
                        column_b["y0_pdf"],
                        column_b["x1"],
                        column_b["y1_pdf"],
                    ],
                    "width": column_b["width"],
                    "height": column_b["height"],
                },
            }
        )

    return matches


if __name__ == "__main__":
    cli()
